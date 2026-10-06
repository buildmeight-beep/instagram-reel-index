"""Build overview.xlsx from raw/posts-all.json and notes.json: a Reels sheet (cover image, title,
what it covers, what you get, tools, keyword, stats, link) and a Tools sheet (each tool, how many
reels name it, their total views).
    python excel.py --out ./nateherkai            (needs: pip install openpyxl pillow)
"""
import argparse, datetime as dt, io, json
from collections import defaultdict
from pathlib import Path
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from PIL import Image

ap = argparse.ArgumentParser(); ap.add_argument("--out", required=True); a = ap.parse_args()
OUT = Path(a.out)
posts = json.loads((OUT / "raw" / "posts-all.json").read_text(encoding="utf-8"))
notes = json.loads((OUT / "notes.json").read_text(encoding="utf-8"))
handle = posts[0].get("user", {}).get("username") or OUT.name

HEAD = Font(bold=True, color="FFFFFF"); FILL = PatternFill("solid", fgColor="1F4FD1")
WRAP = Alignment(wrap_text=True, vertical="top"); TOP = Alignment(vertical="top")

wb = Workbook()
ws = wb.active; ws.title = "Reels"
cols = [("Cover", 10), ("Date", 11), ("Title", 34), ("Title from", 9), ("What it covers", 45), ("What you get", 45),
        ("Tools", 26), ("Keyword", 14), ("Freebie", 24), ("Views", 11), ("Likes", 9), ("Comments", 10),
        ("Length (s)", 9), ("Sponsored", 10), ("Link", 14)]
ws.append([c for c, _ in cols])
for i, (_, w) in enumerate(cols, 1):
    ws.column_dimensions[get_column_letter(i)].width = w
    ws.cell(1, i).font = HEAD; ws.cell(1, i).fill = FILL; ws.cell(1, i).alignment = Alignment(vertical="center")

tools = defaultdict(lambda: [0, 0])
for r, p in enumerate(sorted(posts, key=lambda p: p["taken_at"], reverse=True), start=2):
    n = notes.get(p["code"], {})
    views = p.get("ig_play_count") or p.get("play_count") or 0
    link = f"https://www.instagram.com/reel/{p['code']}/"
    ws.append(["", dt.datetime.fromtimestamp(p["taken_at"], dt.timezone.utc).date(), n.get("title", ""),
               n.get("title_source", ""), n.get("about", ""), n.get("gain", ""), ", ".join(n.get("tools", [])),
               n.get("keyword", ""), n.get("freebie", ""), views, p.get("like_count") or 0,
               p.get("comment_count") or 0, round(p.get("video_duration") or 0),
               "yes" if p.get("is_paid_partnership") else "", "Watch"])
    for c in range(1, len(cols) + 1):
        ws.cell(r, c).alignment = WRAP if c in (3, 5, 6, 7, 9) else TOP
    ws.cell(r, 2).number_format = "yyyy-mm-dd"
    for c in (10, 11, 12):
        ws.cell(r, c).number_format = "#,##0"
    ws.cell(r, 3).font = Font(bold=True)
    ws.cell(r, 15).hyperlink = link; ws.cell(r, 15).style = "Hyperlink"
    ws.row_dimensions[r].height = 80
    thumb = OUT / "thumbs" / f"{p['code']}.jpg"
    if thumb.exists():
        im = Image.open(thumb).convert("RGB"); im.thumbnail((60, 106))
        buf = io.BytesIO(); im.save(buf, "JPEG", quality=80); buf.seek(0)
        xi = XLImage(buf); xi.anchor = f"A{r}"; ws.add_image(xi)
    for t in n.get("tools", []):
        tools[t][0] += 1; tools[t][1] += views
ws.freeze_panes = "C2"
ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{ws.max_row}"

ts = wb.create_sheet("Tools")
ts.append(["Tool", "Reels", "Total views"])
for i, w in enumerate((32, 10, 14), 1):
    ts.column_dimensions[get_column_letter(i)].width = w
    ts.cell(1, i).font = HEAD; ts.cell(1, i).fill = FILL
for t, (cnt, v) in sorted(tools.items(), key=lambda x: (-x[1][0], -x[1][1])):
    ts.append([t, cnt, v]); ts.cell(ts.max_row, 3).number_format = "#,##0"
ts.freeze_panes = "A2"; ts.auto_filter.ref = f"A1:C{ts.max_row}"

wb.properties.title = f"@{handle} reels overview"
wb.save(OUT / "overview.xlsx")
print(f"overview.xlsx: {ws.max_row - 1} reels, {len(tools)} tools")
